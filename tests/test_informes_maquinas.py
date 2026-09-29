"""
Informes de máquinas sin las colas de logística (sep-2026).

Las colas «pedidos esperando documento» y «facturados esperando despacho»
salieron de la página y de los Excel porque gerencia ya no las usa. Los
pedidos de flete se siguen leyendo, pero solo para reconocer un rechazo que el
vendedor volvió a ingresar. Estas pruebas cuidan las dos cosas: que las hojas
de cola no vuelvan y que el seguimiento de rechazos no pierda los pedidos.

RUT ficticios: el repo es público.

    python -m unittest discover tests
"""
import io
import unittest
from datetime import date

import pandas as pd
from openpyxl import load_workbook

from app.export_maquinas import (EN_RUTA, ENTREGADA, RECHAZADA, SIN_DESPACHO,
                                 _prep_pedidos, libro_maquinas,
                                 preparar_movimientos)
from app.export_maquinas_gerencia import libro_gerencia
from app.kpis_maquinas import REINTENTO_PEDIDO, seguimiento_rechazos

GN = 2
HOY = date(2026, 9, 29)
INI, FIN = date(2026, 9, 21), date(2026, 9, 27)

RUT_INST = "11.111.111-1"   # instalación entregada
RUT_RET = "22.222.222-2"    # retiro rechazado, reingresado como pedido
RUT_CAMB = "33.333.333-3"   # cambio sin fila en despachos
RUT_NC = "44.444.444-4"     # NC que anula un retiro

HOJAS_COLA = {"Pedidos sin documento", "Facturados sin despacho",
              "Pendientes de gestionar", "Gestión por vendedor"}


def _maquinas():
    filas = [("5001", "2026-09-22", RUT_INST, "nueva"),
             ("5002", "2026-09-23", RUT_RET, "retiro"),
             ("5003", "2026-09-24", RUT_CAMB, "cambio"),
             ("901", "2026-09-25", RUT_NC, "retiro")]
    return pd.DataFrame([{"documento": d, "fecha": f, "vendedor_id": 7,
                          "cliente_rut": r, "tipo_mov": t, "estado": "gestionada",
                          "sociedad_id": GN} for d, f, r, t in filas])


def _lineas_fl():
    filas = [("5001", RUT_INST, "FL-4", "FACTURA ELECTRONICA"),
             ("5002", RUT_RET, "FL-2", "FACTURA ELECTRONICA"),
             ("5003", RUT_CAMB, "FL-1", "FACTURA ELECTRONICA"),
             ("901", RUT_NC, "FL-2", "NOTA DE CREDITO ELECTRONICA")]
    return pd.DataFrame([{"n_dcto": d, "cliente_rut": r, "producto_codigo": c,
                          "cantidad": 1, "tipo_dcto": t, "sociedad_id": GN}
                         for d, r, c, t in filas])


def _despachos():
    base = {"sociedad_id": GN, "vendedor_id": 7, "transportista": "Cancino Temuco",
            "devolucion": False, "peso": 0, "motivo_rechazo": None}
    return pd.DataFrame([
        dict(base, documento="5001", cliente_rut=RUT_INST, estado="Entregada",
             fecha_ruta="2026-09-23", comentario_entrega=""),
        dict(base, documento="5002", cliente_rut=RUT_RET, estado="Rechazada",
             fecha_ruta="2026-09-24", comentario_entrega="aun con producto"),
        # Una fila de helados del mismo mes: el mes queda con despachos cargados,
        # así que el cambio 5003 sale «Sin despacho» y no «Sin información».
        dict(base, documento="7777", cliente_rut="55.555.555-5",
             estado="Entregada", fecha_ruta="2026-09-24", comentario_entrega=""),
    ])


def _pedidos():
    return pd.DataFrame([{
        "n_pedido": "P-1", "fecha": "2026-09-26", "fecha_pedido": "2026-09-26",
        "producto_codigo": "FL-2", "doc_venta": "Sin DTE", "num_documento": None,
        "estado_pedido": "pending", "cliente_rut": RUT_RET, "vendedor_id": 7,
    }])


def _hojas(xlsx: bytes) -> list:
    return load_workbook(io.BytesIO(xlsx)).sheetnames


class TestInformesSinColas(unittest.TestCase):

    def setUp(self):
        self.mov = preparar_movimientos(_maquinas(), _despachos(), _lineas_fl(),
                                        hoy=HOY)
        self.ped = _prep_pedidos(_pedidos())

    def test_estados_y_nc_aparte(self):
        est = self.mov.set_index("_doc")["Estado entrega"].to_dict()
        self.assertEqual(est["5001"], ENTREGADA)
        self.assertEqual(est["5002"], RECHAZADA)
        self.assertEqual(est["5003"], SIN_DESPACHO)
        self.assertTrue(self.mov.set_index("_doc").loc["901", "_nc"])
        self.assertNotIn(EN_RUTA, est.values())

    def test_rechazo_reingresado_se_sigue_reconociendo(self):
        mov = self.mov[~self.mov["_nc"]]
        seg = seguimiento_rechazos(mov[mov["Estado entrega"] == RECHAZADA],
                                   mov, self.ped, hoy=HOY)
        self.assertEqual(seg.iloc[0]["Estado post-rechazo"], REINTENTO_PEDIDO)
        self.assertFalse(bool(seg.iloc[0]["_abierto"]))

    def test_excel_de_gerencia_sin_hojas_de_cola(self):
        metas = {"meta_gestiones_semana": 22, "meta_pct_entregado": 0.85}
        hojas = _hojas(libro_gerencia(self.mov, self.ped, INI, FIN, metas,
                                      hoy=HOY))
        self.assertFalse(HOJAS_COLA & set(hojas), hojas)
        for h in ("Resumen", "Rechazos · seguimiento", "Sin confirmar",
                  "Fletes anulados (NC)"):
            self.assertIn(h, hojas)

    def test_informe_de_analisis_sin_hojas_de_cola(self):
        hojas = _hojas(libro_maquinas(_maquinas(), INI, FIN,
                                      despachos=_despachos(),
                                      lineas_fl=_lineas_fl(),
                                      pedidos_fl=_pedidos(), hoy=HOY))
        self.assertFalse(HOJAS_COLA & set(hojas), hojas)
        for h in ("Resumen", "Semanal", "Pendientes por confirmar",
                  "Conciliación Autoventa", "Definiciones"):
            self.assertIn(h, hojas)


if __name__ == "__main__":
    unittest.main()
