"""
Reporte semanal para el área comercial (Drive): lo que comercial escribe en las
columnas amarillas tiene que sobrevivir a la regeneración del lunes siguiente,
aunque el caso cambie de lista, y lo que se resolvió tiene que pasar a «Casos
cerrados» con cómo se cerró.

RUT ficticios: el repo es público.

    python -m unittest discover tests
"""
import io
import unittest
from datetime import date

import pandas as pd
from openpyxl import load_workbook

from app.export_maquinas import RECHAZADA, _prep_pedidos, preparar_movimientos
from app.export_seguimiento_comercial import (
    COL_ESTADO, HOJA_CERRADOS, HOJA_HISTORIAL, HOJA_PEDIDOS, HOJA_REGULARIZADAS,
    HOJA_SIN_CERRAR, HOJA_SIN_RETOMAR, ID, id_rechazo, libro_seguimiento)
from app.kpis_maquinas import REGULARIZADA, resumen_rechazos, seguimiento_rechazos

GN = 2
RUT_A = "11.111.111-1"   # retiro rechazado, después reingresado
RUT_B = "22.222.222-2"   # pedido de instalación sin DTE


def _mov(con_reintento: bool, hoy):
    filas = [("5002", "2026-09-23", RUT_A, "retiro")]
    if con_reintento:
        filas.append(("5100", "2026-10-08", RUT_A, "retiro"))
    maq = pd.DataFrame([{"documento": d, "fecha": f, "vendedor_id": 7,
                         "cliente_rut": r, "tipo_mov": t, "estado": "gestionada",
                         "sociedad_id": GN} for d, f, r, t in filas])
    desp = pd.DataFrame([
        dict(documento="5002", cliente_rut=RUT_A, estado="Rechazada",
             fecha_ruta="2026-09-24", transportista="Máquinas Kreems 1",
             comentario_entrega="local cerrado", sociedad_id=GN, vendedor_id=7,
             devolucion=False, peso=0, motivo_rechazo=None),
        dict(documento="5100", cliente_rut=RUT_A, estado="Pendiente",
             fecha_ruta="2026-10-09", transportista="Máquinas Kreems 1",
             comentario_entrega="", sociedad_id=GN, vendedor_id=7,
             devolucion=False, peso=0, motivo_rechazo=None),
    ])
    return preparar_movimientos(maq, desp, hoy=hoy)


def _ped(facturado: bool):
    return _prep_pedidos(pd.DataFrame([{
        "n_pedido": "6001", "fecha": "2026-10-02", "fecha_pedido": "2026-10-01",
        "producto_codigo": "FL-4",
        "doc_venta": "Factura" if facturado else "Sin DTE",
        "num_documento": "5200" if facturado else None,
        "estado_pedido": "invoiced" if facturado else "pending",
        "cliente_rut": RUT_B, "vendedor_id": 7}]))


def _mq(fecha="2026-10-10"):
    """Una factura MQ («máquinas regularizadas») del cliente del rechazo."""
    return pd.DataFrame([{
        "cliente_rut": RUT_A, "_rutn": "111111111", "documento": "6667",
        "fecha": pd.Timestamp(fecha), "comentario_mq": "máquina retirada, mal cerrada",
        "estado_despacho_mq": "Entregada"}])


def _generar(con_reintento, facturado, hoy, previo=None, mq=None):
    mov = _mov(con_reintento, hoy)
    ped = _ped(facturado)
    seg = seguimiento_rechazos(mov[mov["Estado entrega"] == RECHAZADA], mov,
                               ped, hoy=hoy, mq=mq)
    return libro_seguimiento(seg, ped, previo=previo, vendedores={7: "Ana"},
                             hoy=hoy)


def _tabla(xlsx, hoja):
    ws = load_workbook(io.BytesIO(xlsx))[hoja]
    filas = list(ws.iter_rows(values_only=True))
    i = next(k for k, f in enumerate(filas) if ID in f or "Fecha" in f)
    cols = filas[i]
    return [dict(zip(cols, f)) for f in filas[i + 1:] if f[0] is not None]


def _escribir_comentario(xlsx, hoja, id_caso, columna, valor) -> bytes:
    wb = load_workbook(io.BytesIO(xlsx))
    ws = wb[hoja]
    hdr = [c.value for c in ws[3]]
    for row in ws.iter_rows(min_row=4):
        if row[0].value == id_caso:
            row[hdr.index(columna)].value = valor
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestSeguimientoComercial(unittest.TestCase):

    def test_semana_1_las_tres_listas(self):
        xlsx, res = _generar(False, False, date(2026, 10, 5))
        self.assertEqual(res["sin_retomar"], 1)
        self.assertEqual(res["sin_cerrar"], 0)
        self.assertEqual(res["pedidos_sin_gestionar"], 1)
        fila = _tabla(xlsx, HOJA_SIN_RETOMAR)[0]
        self.assertEqual(fila[ID], id_rechazo("5002", "2026-09-23"))
        self.assertEqual(fila[COL_ESTADO], "Pendiente")
        self.assertEqual(_tabla(xlsx, HOJA_PEDIDOS)[0][ID], "P-6001")

    def test_comentarios_sobreviven_y_el_caso_cambia_de_lista(self):
        rid = id_rechazo("5002", "2026-09-23")
        x1, _ = _generar(False, False, date(2026, 10, 5))
        x1 = _escribir_comentario(x1, HOJA_SIN_RETOMAR, rid,
                                  "Comentario / situación del cliente",
                                  "Ana lo reingresa el jueves")
        x1 = _escribir_comentario(x1, HOJA_SIN_RETOMAR, rid, COL_ESTADO,
                                  "En gestión")
        x1 = _escribir_comentario(x1, HOJA_PEDIDOS, "P-6001", COL_ESTADO,
                                  "En gestión")
        # Semana 2: el retiro se reingresó (doc 5100 en ruta) y el pedido se
        # facturó.
        x2, res = _generar(True, True, date(2026, 10, 12), previo=x1)
        self.assertTrue(res["previo_leido"])
        self.assertEqual(res["sin_retomar"], 0)
        self.assertEqual(res["sin_cerrar"], 1)
        fila = _tabla(x2, HOJA_SIN_CERRAR)[0]
        self.assertEqual(fila[ID], rid)
        self.assertEqual(fila["Comentario / situación del cliente"],
                         "Ana lo reingresa el jueves")
        self.assertEqual(fila[COL_ESTADO], "En gestión")
        self.assertEqual(fila["Semanas en el reporte"], 1)
        cerr = _tabla(x2, HOJA_CERRADOS)
        self.assertEqual([c[ID] for c in cerr], ["P-6001"])
        self.assertIn("DTE 5200", cerr[0]["Cómo se cerró"])
        self.assertEqual(cerr[0][COL_ESTADO], "En gestión")
        self.assertEqual(len(_tabla(x2, HOJA_HISTORIAL)), 2)

    def test_cerrados_se_acumulan_y_regenerar_el_mismo_dia_no_duplica(self):
        x1, _ = _generar(False, False, date(2026, 10, 5))
        x2, _ = _generar(True, True, date(2026, 10, 12), previo=x1)
        x3, res = _generar(True, True, date(2026, 10, 12), previo=x2)
        self.assertEqual([c[ID] for c in _tabla(x3, HOJA_CERRADOS)], ["P-6001"])
        self.assertEqual(len(_tabla(x3, HOJA_HISTORIAL)), 2)

    def test_mq_cierra_el_rechazo_y_queda_trazado(self):
        rid = id_rechazo("5002", "2026-09-23")
        x1, _ = _generar(False, False, date(2026, 10, 5))
        # Semana 2: nadie lo reingresó, pero se emitió una factura MQ.
        x2, res = _generar(False, False, date(2026, 10, 12), previo=x1, mq=_mq())
        self.assertEqual(res["sin_retomar"], 0)
        self.assertEqual(res["sin_cerrar"], 0)
        self.assertEqual(res["regularizadas"], 1)
        cerr = [c for c in _tabla(x2, HOJA_CERRADOS) if c[ID] == rid]
        self.assertEqual(len(cerr), 1)
        self.assertIn("Regularizada con MQ 6667 del 10/10", cerr[0]["Cómo se cerró"])
        self.assertIn("«máquina retirada, mal cerrada»", cerr[0]["Cómo se cerró"])
        self.assertIn("«local cerrado»", cerr[0]["Cómo se cerró"])
        reg = _tabla(x2, HOJA_REGULARIZADAS)[0]
        self.assertEqual(reg[ID], rid)
        self.assertEqual(reg["Estado antes de regularizar"], "Sin reintento")
        self.assertEqual(reg["Factura MQ"], "6667")
        self.assertEqual(reg["Comentario del rechazo"], "local cerrado")
        self.assertEqual(reg["Comentario del despacho MQ"], "máquina retirada, mal cerrada")

    def test_mq_anterior_al_rechazo_no_cierra(self):
        hoy = date(2026, 10, 12)
        mov = _mov(False, hoy)
        seg = seguimiento_rechazos(mov[mov["Estado entrega"] == RECHAZADA], mov,
                                   _ped(False), hoy=hoy, mq=_mq("2026-09-01"))
        self.assertNotIn(REGULARIZADA, set(seg["Estado post-rechazo"]))
        self.assertEqual(resumen_rechazos(seg)["regularizadas"], 0)

    def test_previo_ilegible_se_informa(self):
        _, res = _generar(False, False, date(2026, 10, 5), previo=b"no es xlsx")
        self.assertFalse(res["previo_leido"])


if __name__ == "__main__":
    unittest.main()
