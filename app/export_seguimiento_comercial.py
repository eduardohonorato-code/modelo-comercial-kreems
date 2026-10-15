"""
Reporte semanal para el área comercial: los casos de máquinas que alguien
tiene que mover, con columnas para que comercial anote qué hizo.

Nace de la reunión de los lunes (oct-2026). Control de Máquinas mostraba los
rechazos, pero la lista se quedaba en la pantalla de gerencia: nadie se la
mandaba al vendedor y, a la semana siguiente, no había cómo saber qué se hizo
con cada uno. Este libro se sube a Drive después de cada reunión y comercial lo
trabaja ahí. Tres listas, una por tipo de pendiente:

  1. Rechazos sin retomar      · volvieron al camión y nadie los reingresó.
                                 Responsable: el vendedor.
  2. Reingresos sin cerrar     · se volvieron a ingresar y aún no terminan
                                 (en ruta, esperando DTE, rechazados de nuevo).
  3. Pedidos sin gestionar     · pedidos de flete ingresados que todavía no
                                 tienen documento. Responsable: logística.

LA REGLA QUE HACE QUE SIRVA: las columnas amarillas son de comercial y se
CONSERVAN de una semana a otra. Cada caso tiene un «ID caso» estable
(R-<documento>-<AAAAMMDD> para un rechazo, P-<n° pedido> para un pedido), y al
regenerar el libro se le pasa la versión anterior: lo que comercial escribió se
copia a la fila del mismo ID, aunque el caso haya cambiado de lista (un rechazo
que se reingresó pasa de la 1 a la 2 con sus comentarios). Lo que estaba
abierto y ya no está pasa a «4 · Casos cerrados» con cómo se cerró, y queda ahí
como historial. «Historial semanal» guarda los conteos de cada reunión.

El cálculo del seguimiento NO vive aquí: viene hecho de
`kpis_maquinas.seguimiento_rechazos` (el mismo que usa la página), para que el
reporte y la pantalla nunca digan cosas distintas.
"""
import io
from datetime import date

import pandas as pd

from app.kpis_maquinas import (ANULADA, DIAS_PARA_REINTENTAR, LECT_DEJO,
                               LECT_IGUAL, LECT_MENOS, LECT_NO, LECT_NUEVO,
                               LECT_RECIENTE, LECT_SIN_BASE, REINTENTO_OK,
                               REINTENTO_OTRO, REINTENTO_PEDIDO,
                               REINTENTO_RECHAZO, REINTENTO_RUTA,
                               REINTENTO_SININFO)

_MOV = {"nueva": "Instalación", "cambio": "Cambio", "retiro": "Retiro"}

HOJA_RESUMEN = "Resumen"
HOJA_SIN_RETOMAR = "1 · Rechazos sin retomar"
HOJA_SIN_CERRAR = "2 · Reingresos sin cerrar"
HOJA_PEDIDOS = "3 · Pedidos sin gestionar"
HOJA_CERRADOS = "4 · Casos cerrados"
HOJA_HISTORIAL = "Historial semanal"
HOJAS_CASOS = (HOJA_SIN_RETOMAR, HOJA_SIN_CERRAR, HOJA_PEDIDOS)

_TIPO_CASO = {HOJA_SIN_RETOMAR: "Rechazo sin retomar",
              HOJA_SIN_CERRAR: "Reingreso sin cerrar",
              HOJA_PEDIDOS: "Pedido sin gestionar"}

ID = "ID caso"
DESDE = "En el reporte desde"
SEMANAS = "Semanas en el reporte"

# Las columnas que llena comercial. Son las únicas que se copian de la versión
# anterior; todo lo demás se recalcula desde la base cada lunes.
COL_ESTADO = "Estado gestión"
COLS_MANUALES = [COL_ESTADO, "Acción comprometida", "Responsable comercial",
                 "Fecha compromiso", "Comentario / situación del cliente"]
ESTADOS_GESTION = ["Pendiente", "En gestión", "Reingresado", "Descartado",
                   "Cerrado"]

_HDR_BG = "1B3A6B"       # navy de la casa
_MANUAL_BG = "FFF4C2"    # amarillo: lo llena comercial
_MANUAL_HDR = "B7791F"   # encabezado de las columnas manuales
_FMT_FECHA = "dd/mm/yyyy"
_FMT_CLP = '"$"#,##0;[Red]-"$"#,##0'

# Quién tiene que mover cada estado. «El pedido es del vendedor; emitir el DTE
# y despachar son de logística» (decisión de gerencia, sep-2026).
_RESPONSABLE = {
    REINTENTO_RUTA: "Logística (confirmar entrega)",
    REINTENTO_SININFO: "Logística (confirmar entrega)",
    REINTENTO_PEDIDO: "Logística (emitir DTE)",
    REINTENTO_RECHAZO: "Vendedor",
    REINTENTO_OTRO: "Vendedor (confirmar)",
}
_ACCION = {
    REINTENTO_RUTA: "Confirmar con logística que se entregó.",
    REINTENTO_SININFO: "Confirmar con logística que se entregó.",
    REINTENTO_PEDIDO: "Logística debe emitir el DTE del pedido reingresado.",
    REINTENTO_RECHAZO: "Segundo rechazo: hablar con el cliente antes de "
                       "volver a ingresarlo.",
    REINTENTO_OTRO: "El cliente tuvo otro movimiento: confirmar si reemplaza "
                    "al rechazado y, si es así, cerrar.",
}


# ── Las tres listas ──────────────────────────────────────────────────────────

def id_rechazo(documento, fecha_documento) -> str:
    """ID estable de un rechazo. El número de documento se reutiliza entre
    meses, así que va con la fecha del documento."""
    f = pd.Timestamp(fecha_documento)
    return f"R-{documento}-{f:%Y%m%d}" if pd.notna(f) else f"R-{documento}"


def id_pedido(n_pedido) -> str:
    return f"P-{n_pedido}"


def _accion_sin_retomar(lectura) -> str:
    """Qué pedirle al vendedor según lo que el cliente ha comprado después."""
    if lectura in (LECT_IGUAL, LECT_SIN_BASE, LECT_NUEVO):
        return ("Sigue comprando: confirmar si se renegoció. Si el movimiento "
                "ya no aplica, descartarlo con el motivo.")
    if lectura in (LECT_NO, LECT_DEJO, LECT_MENOS):
        return ("Visitar al cliente: reingresar el pedido o descartarlo con "
                "el motivo.")
    if lectura == LECT_RECIENTE:
        return "Rechazo reciente: reingresar el pedido esta semana."
    return "Reingresar el pedido o informar por qué no se hará."


def casos_abiertos(seg: pd.DataFrame, ped: pd.DataFrame,
                   ventas_post: pd.DataFrame | None = None,
                   vendedores: dict | None = None,
                   clientes: pd.DataFrame | None = None,
                   hoy: date | None = None) -> dict:
    """
    Las tres listas, como DataFrames listos para escribir (sin las columnas
    manuales ni las de antigüedad: esas las pone `libro_seguimiento`).

    `seg` es `seguimiento_rechazos` sobre TODOS los rechazos desde el primer
    despacho cargado (no solo los de la semana): un rechazo de hace dos meses
    que nadie retomó sigue siendo trabajo pendiente.
    `ped` es `_prep_pedidos` de los pedidos FL. `ventas_post` es
    `ventas_post_rechazo` de los sin retomar (opcional).
    """
    hoy_ts = pd.Timestamp(hoy or date.today())
    out = {}
    s = seg if seg is not None else pd.DataFrame()

    # 1 · Sin retomar
    if s.empty:
        sr = pd.DataFrame()
    else:
        g = s[s["_abierto"]].copy()
        if ventas_post is not None and not ventas_post.empty:
            g = g.join(ventas_post[["Lectura de la venta", "Última compra"]])
        else:
            g["Lectura de la venta"] = None
            g["Última compra"] = None
        sr = pd.DataFrame({
            ID: [id_rechazo(d, f) for d, f in zip(g["Documento"],
                                                  g["Fecha documento"])],
            "Vendedor": g["Vendedor"].values,
            "Cliente": g["Cliente"].values,
            "RUT": g["RUT"].values,
            "Comuna": g["Comuna"].values,
            "Movimiento": g["Movimiento"].values,
            "Fecha rechazo": g["Fecha rechazo"].values,
            "Días desde el rechazo": g["Días desde el rechazo"].values,
            "Motivo": g["Motivo"].values,
            "Lo que dijo el repartidor": g["Lo que dijo el repartidor"].values,
            "Lectura de la venta": g["Lectura de la venta"].fillna("—").values,
            "Última compra": g["Última compra"].values,
            "Responsable": "Vendedor",
            "Acción sugerida": [_accion_sin_retomar(x)
                                for x in g["Lectura de la venta"]],
            "Documento rechazado": g["Documento"].values,
        }).sort_values(["Vendedor", "Días desde el rechazo"],
                       ascending=[True, False])
    out[HOJA_SIN_RETOMAR] = sr

    # 2 · Reingresados sin cerrar (lo que la página llama «sin cerrar»)
    if s.empty:
        sc = pd.DataFrame()
    else:
        est = s["Estado post-rechazo"]
        g = s[~s["_abierto"] & ~est.isin([REINTENTO_OK, ANULADA])].copy()
        sc = pd.DataFrame({
            ID: [id_rechazo(d, f) for d, f in zip(g["Documento"],
                                                  g["Fecha documento"])],
            "Vendedor": g["Vendedor"].values,
            "Cliente": g["Cliente"].values,
            "RUT": g["RUT"].values,
            "Comuna": g["Comuna"].values,
            "Movimiento": g["Movimiento"].values,
            "Fecha rechazo": g["Fecha rechazo"].values,
            "Días desde el rechazo": g["Días desde el rechazo"].values,
            "Motivo": g["Motivo"].values,
            "Estado post-rechazo": g["Estado post-rechazo"].values,
            "Reintento": g["Reintento"].values,
            "Responsable": g["Estado post-rechazo"].map(_RESPONSABLE)
                           .fillna("Vendedor").values,
            "Acción sugerida": g["Estado post-rechazo"].map(_ACCION)
                               .fillna("Revisar el caso.").values,
            "Documento rechazado": g["Documento"].values,
        }).sort_values(["Responsable", "Días desde el rechazo"],
                       ascending=[True, False])
    out[HOJA_SIN_CERRAR] = sc

    # 3 · Pedidos sin gestionar (sin DTE, descartados los fantasma y los
    # facturados en la otra serie — ver _prep_pedidos)
    cola = (ped[ped["_sin_dte"] & ~ped["_fantasma"]].copy()
            if ped is not None and not ped.empty else pd.DataFrame())
    if cola.empty:
        pg = pd.DataFrame()
    else:
        # Un pedido puede tener varias líneas FL (dos máquinas): una fila por pedido.
        cola = cola.sort_values("_ingreso").drop_duplicates("n_pedido")
        nom = _nombres(clientes)
        com = _nombres(clientes, "comuna")
        reing = set()
        if not s.empty:
            # «pedido 6418 del 09/10, sin DTE hace 6 días» → 6418
            txt = s.loc[s["Estado post-rechazo"] == REINTENTO_PEDIDO, "Reintento"]
            reing = set(txt.astype(str).str.extract(r"pedido (\S+)")[0].dropna())
        vmap = vendedores or {}
        pg = pd.DataFrame({
            ID: [id_pedido(n) for n in cola["n_pedido"]],
            "Vendedor": cola["vendedor_id"].map(vmap).fillna("Sin asignar").values,
            "Cliente": cola["cliente_rut"].map(nom).fillna(cola["cliente_rut"]).values,
            "RUT": cola["cliente_rut"].values,
            "Comuna": cola["cliente_rut"].map(com).fillna("(sin dato)").values,
            "Movimiento": cola["_mov"].map(_MOV).fillna("(otro)").values,
            "Fecha pedido": cola["_ingreso"].dt.date.values,
            "Días esperando": (hoy_ts - cola["_ingreso"]).dt.days.values,
            "N° pedido": cola["n_pedido"].values,
            "Es reingreso de un rechazo": ["Sí" if str(n) in reing else "No"
                                           for n in cola["n_pedido"]],
            "Responsable": "Logística (emitir DTE)",
            "Acción sugerida": "Emitir el DTE y programar la ruta, o avisar "
                               "al vendedor por qué no sale.",
        }).sort_values("Días esperando", ascending=False)
    out[HOJA_PEDIDOS] = pg
    return out


def _nombres(clientes: pd.DataFrame | None, col: str = "razon_social") -> dict:
    if clientes is None or clientes.empty or col not in clientes.columns:
        return {}
    c = clientes.drop_duplicates("rut")
    return dict(zip(c["rut"], c[col].astype(str).str.strip()))


def como_se_cerro(id_caso: str, seg: pd.DataFrame, ped: pd.DataFrame) -> str:
    """Por qué un caso que estaba abierto la semana pasada ya no lo está."""
    if id_caso.startswith("R-") and seg is not None and not seg.empty:
        ids = [id_rechazo(d, f) for d, f in zip(seg["Documento"],
                                               seg["Fecha documento"])]
        m = seg[[i == id_caso for i in ids]]
        if not m.empty:
            r = m.iloc[0]
            return f"{r['Estado post-rechazo']} ({r['Reintento']})"
        return "Ya no aparece en el seguimiento (fuera de la ventana cargada)"
    if id_caso.startswith("P-") and ped is not None and not ped.empty:
        n = id_caso[2:]
        m = ped[ped["n_pedido"].astype(str) == n]
        if m.empty:
            return "El pedido ya no existe en la base"
        r = m.iloc[0]
        if bool(r.get("_facturado_otra_serie")):
            return "Facturado en la otra serie (Obuma)"
        if bool(r.get("_fantasma")):
            return "Anulado o reingresado con otro número en Autoventa"
        if pd.notna(r.get("num_documento")):
            f = r.get("fecha")
            return (f"DTE {r['num_documento']} emitido"
                    + (f" el {pd.Timestamp(f):%d/%m/%Y}" if pd.notna(f) else ""))
    return "Cerrado"


def datos_seguimiento(client, hoy: date | None = None) -> dict:
    """
    Todo lo que necesita el libro, leído de la base. Lo usa el script de los
    lunes (reportes/seguimiento_comercial_maquinas.py); la página ya tiene
    estos datos cargados y llama a `libro_seguimiento` directo.

    Misma ventana que Control de Máquinas: desde un mes antes del primer
    despacho hasta 14 días después de hoy (logística factura con la fecha de
    entrega), y sin las NC de flete, que anulan y no son gestiones.
    """
    from datetime import timedelta

    from app.data import get_todos_vendedores, get_dim_cliente_full, \
        get_ventas_clientes
    from app.export_maquinas import RECHAZADA
    from app.kpis_maquinas import (PRIMER_DESPACHO, cargar_todo, desde_ventas,
                                   seguimiento_rechazos, ventas_post_rechazo)

    hoy = hoy or date.today()
    mov, ped, _ = cargar_todo(client, PRIMER_DESPACHO - timedelta(days=31),
                              hoy + timedelta(days=14))
    if mov is not None and not mov.empty and "_nc" in mov.columns:
        mov = mov[~mov["_nc"]]
    seg = (seguimiento_rechazos(mov[mov["Estado entrega"] == RECHAZADA], mov,
                                ped, hoy=hoy)
           if mov is not None and not mov.empty else pd.DataFrame())
    vp = None
    if not seg.empty and seg["_abierto"].any():
        sin = seg[seg["_abierto"]]
        try:
            ventas = get_ventas_clientes(client, sorted(set(sin["RUT"])),
                                         desde_ventas(sin))
            vp = ventas_post_rechazo(sin, ventas, hoy)
        except Exception:
            vp = None
    try:
        v = get_todos_vendedores(client)
        vend = dict(zip(v["id"], v["nombre_canonico"]))
    except Exception:
        vend = {}
    try:
        cli = get_dim_cliente_full(client)
    except Exception:
        cli = None
    return dict(seg=seg, ped=ped, ventas_post=vp, vendedores=vend,
                clientes=cli)


# ── Lectura de la versión anterior ───────────────────────────────────────────

def _encabezado(ws, buscar: str = ID, max_filas: int = 8):
    """Fila del encabezado (la que contiene `buscar`) y sus columnas."""
    for fila in range(1, max_filas + 1):
        vals = [c.value for c in ws[fila]]
        if buscar in vals:
            return fila, {str(v): j for j, v in enumerate(vals) if v is not None}
    return None, {}


def _filas(ws, buscar: str = ID) -> list[dict]:
    fila0, cols = _encabezado(ws, buscar)
    if not fila0:
        return []
    out = []
    for row in ws.iter_rows(min_row=fila0 + 1, values_only=True):
        d = {k: row[j] if j < len(row) else None for k, j in cols.items()}
        if d.get(buscar) in (None, ""):
            continue
        out.append(d)
    return out


def leer_previo(previo) -> dict:
    """
    Lo que hay que rescatar del libro de la semana pasada.

    `previo` son los bytes del xlsx, una ruta, o None. Devuelve:
      manual   · {ID: {columna manual: valor}}
      desde    · {ID: fecha en que entró al reporte}
      abiertos · {ID: fila completa + "Tipo de caso"} de las tres listas
      cerrados · las filas de «4 · Casos cerrados» (para no perder historia)
      historial· las filas de «Historial semanal»
    Si el archivo no se puede leer, se devuelve todo vacío: mejor un reporte
    sin comentarios que ningún reporte (y el llamador lo avisa).
    """
    vacio = dict(manual={}, desde={}, abiertos={}, cerrados=[], historial=[])
    if previo is None:
        return vacio
    from openpyxl import load_workbook
    try:
        src = io.BytesIO(previo) if isinstance(previo, (bytes, bytearray)) else previo
        wb = load_workbook(src, data_only=True)
    except Exception:
        return vacio
    res = dict(manual={}, desde={}, abiertos={}, cerrados=[], historial=[])
    for hoja in HOJAS_CASOS:
        if hoja not in wb.sheetnames:
            continue
        for d in _filas(wb[hoja]):
            i = str(d[ID])
            res["manual"][i] = {c: d.get(c) for c in COLS_MANUALES}
            res["desde"][i] = _a_fecha(d.get(DESDE))
            res["abiertos"][i] = dict(d, **{"Tipo de caso": _TIPO_CASO[hoja]})
    if HOJA_CERRADOS in wb.sheetnames:
        res["cerrados"] = _filas(wb[HOJA_CERRADOS])
    if HOJA_HISTORIAL in wb.sheetnames:
        res["historial"] = _filas(wb[HOJA_HISTORIAL], "Fecha")
    return res


def _a_fecha(v):
    if v is None or v == "":
        return None
    try:
        return pd.Timestamp(v).date()
    except (ValueError, TypeError):
        return None


# ── El libro ─────────────────────────────────────────────────────────────────

def libro_seguimiento(seg: pd.DataFrame, ped: pd.DataFrame,
                      previo=None, ventas_post: pd.DataFrame | None = None,
                      vendedores: dict | None = None,
                      clientes: pd.DataFrame | None = None,
                      hoy: date | None = None) -> tuple[bytes, dict]:
    """
    Arma el libro y devuelve `(xlsx, resumen)`.

    `resumen` trae los conteos (para mostrarlos en la app o en consola) y
    `previo_leido`: False si se pasó una versión anterior que no se pudo leer
    —en ese caso los comentarios de comercial NO se conservaron y hay que
    avisar antes de subir el archivo a Drive.
    """
    from openpyxl import Workbook

    hoy = hoy or date.today()
    prev = leer_previo(previo)
    previo_leido = previo is None or bool(prev["abiertos"] or prev["historial"]
                                          or prev["cerrados"])
    listas = casos_abiertos(seg, ped, ventas_post, vendedores, clientes, hoy)

    # Antigüedad en el reporte + columnas manuales rescatadas por ID
    abiertos_hoy = set()
    for hoja, df in listas.items():
        if df.empty:
            continue
        ids = df[ID].astype(str)
        abiertos_hoy |= set(ids)
        desde = [prev["desde"].get(i) or hoy for i in ids]
        df[DESDE] = desde
        df[SEMANAS] = [(hoy - d).days // 7 for d in desde]
        for c in COLS_MANUALES:
            df[c] = [prev["manual"].get(i, {}).get(c) for i in ids]
        df[COL_ESTADO] = df[COL_ESTADO].fillna("Pendiente")

    # Lo que estaba abierto y ya no está → cerrado esta semana
    nuevos_cerrados = []
    for i, fila in prev["abiertos"].items():
        if i in abiertos_hoy:
            continue
        nuevos_cerrados.append({
            ID: i, "Tipo de caso": fila.get("Tipo de caso"),
            "Vendedor": fila.get("Vendedor"), "Cliente": fila.get("Cliente"),
            "Movimiento": fila.get("Movimiento"),
            DESDE: _a_fecha(fila.get(DESDE)),
            "Cerrado detectado el": hoy,
            "Cómo se cerró": como_se_cerro(i, seg, ped),
            **{c: fila.get(c) for c in COLS_MANUALES},
        })
    ya = {str(r[ID]) for r in nuevos_cerrados}
    # Un caso que se cerró y volvió a abrirse sale de la lista de cerrados.
    viejos = [r for r in prev["cerrados"]
              if str(r[ID]) not in ya and str(r[ID]) not in abiertos_hoy]
    cerrados = pd.DataFrame(nuevos_cerrados + viejos)
    if not cerrados.empty:
        for c in (DESDE, "Cerrado detectado el"):
            cerrados[c] = cerrados[c].map(_a_fecha)
        cerrados = cerrados.sort_values("Cerrado detectado el", ascending=False)

    sr, sc, pg = (listas[HOJA_SIN_RETOMAR], listas[HOJA_SIN_CERRAR],
                  listas[HOJA_PEDIDOS])
    venc = (int((sr["Días desde el rechazo"] > DIAS_PARA_REINTENTAR).sum())
            if not sr.empty else 0)
    resumen = {
        "sin_retomar": len(sr), "sin_retomar_vencidos": venc,
        "sin_cerrar": len(sc), "pedidos_sin_gestionar": len(pg),
        "cerrados_semana": len(nuevos_cerrados),
        "previo_leido": previo_leido,
    }

    # Historial: una fila por reunión (si se regenera el mismo día, se pisa)
    hist = [h for h in prev["historial"] if _a_fecha(h.get("Fecha")) != hoy]
    hist.append({"Fecha": hoy,
                 "Rechazos sin retomar": resumen["sin_retomar"],
                 f"De ellos, más de {DIAS_PARA_REINTENTAR} días": venc,
                 "Reingresos sin cerrar": resumen["sin_cerrar"],
                 "Pedidos sin gestionar": resumen["pedidos_sin_gestionar"],
                 "Cerrados desde el reporte anterior": resumen["cerrados_semana"]})
    historial = pd.DataFrame(hist)
    historial["Fecha"] = historial["Fecha"].map(_a_fecha)
    historial = historial.sort_values("Fecha")

    wb = Workbook()
    wb.remove(wb.active)
    _hoja_resumen(wb, listas, resumen, hoy)
    notas = {
        HOJA_SIN_RETOMAR: "Rechazos que volvieron al camión y que NADIE volvió "
                          "a ingresar. El vendedor debe reingresar el pedido o "
                          "informar por qué no (descartarlo con el motivo).",
        HOJA_SIN_CERRAR: "Rechazos que SÍ se volvieron a ingresar pero todavía "
                         "no terminan. «Responsable» dice quién tiene la pelota.",
        HOJA_PEDIDOS: "Pedidos de flete de máquina que el vendedor ya ingresó "
                      "en Autoventa y que todavía no tienen documento (DTE). "
                      "Todavía no son gestiones.",
    }
    for hoja in HOJAS_CASOS:
        _hoja_casos(wb, hoja, listas[hoja], notas[hoja])
    _hoja_casos(wb, HOJA_CERRADOS, cerrados,
                "Casos que estaban abiertos en un reporte anterior y ya no lo "
                "están. Se acumulan: es el historial de lo que se resolvió.",
                validar=False)
    _hoja_casos(wb, HOJA_HISTORIAL, historial,
                "Conteo de cada reporte. Si los números no bajan, los casos no "
                "se están moviendo.", manuales=False, clave="Fecha")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), resumen


def _hoja_resumen(wb, listas: dict, res: dict, hoy: date) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment

    ws = wb.create_sheet(HOJA_RESUMEN)
    ws.column_dimensions["A"].width = 44
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16
    ws["A1"] = "Seguimiento comercial · Máquinas en comodato"
    ws["A1"].font = Font(bold=True, size=14, color=_HDR_BG)
    ws["A2"] = f"Corte al {hoy:%d/%m/%Y} (reunión de gerencia de los lunes)"
    ws["A2"].font = Font(italic=True, color="64748B")

    filas = [
        ("1 · Rechazos sin retomar", res["sin_retomar"],
         f"{res['sin_retomar_vencidos']} llevan más de "
         f"{DIAS_PARA_REINTENTAR} días"),
        ("2 · Reingresos sin cerrar", res["sin_cerrar"],
         "se volvieron a ingresar y aún no terminan"),
        ("3 · Pedidos sin gestionar", res["pedidos_sin_gestionar"],
         "pedidos de flete sin documento"),
        ("Cerrados desde el reporte anterior", res["cerrados_semana"],
         "ver «4 · Casos cerrados»"),
    ]
    for k, (a, b, c) in enumerate(filas, 4):
        ws.cell(row=k, column=1, value=a).font = Font(bold=True)
        ws.cell(row=k, column=2, value=b).font = Font(bold=True, size=12)
        ws.cell(row=k, column=3, value=c).font = Font(color="64748B")

    # Por vendedor: es la tabla que se reenvía a cada uno.
    tabla = _por_vendedor(listas)
    f0 = 10
    ws.cell(row=f0 - 1, column=1, value="Casos abiertos por vendedor").font = \
        Font(bold=True, color=_HDR_BG)
    if tabla.empty:
        ws.cell(row=f0, column=1, value="Sin casos abiertos.")
    else:
        for j, c in enumerate(tabla.columns, 1):
            x = ws.cell(row=f0, column=j, value=c)
            x.font = Font(bold=True, color="FFFFFF")
            x.fill = PatternFill("solid", fgColor=_HDR_BG)
            x.alignment = Alignment(horizontal="center", wrap_text=True)
        for i, row in enumerate(tabla.itertuples(index=False), f0 + 1):
            for j, v in enumerate(row, 1):
                ws.cell(row=i, column=j, value=int(v) if j > 1 else v)
        ultima = ws.cell(row=f0 + len(tabla), column=1)
        if ultima.value == "TOTAL":
            for j in range(1, len(tabla.columns) + 1):
                ws.cell(row=f0 + len(tabla), column=j).font = Font(bold=True)

    f1 = f0 + len(tabla) + 3
    ws.cell(row=f1, column=1, value="Cómo se usa").font = Font(bold=True,
                                                             color=_HDR_BG)
    pasos = [
        "1. Cada lunes, después de la reunión, gerencia actualiza este archivo.",
        "2. Comercial trabaja SOLO las columnas amarillas (Estado gestión, "
        "acción, responsable, fecha compromiso, comentario).",
        "3. Lo escrito en amarillo se conserva la semana siguiente, aunque el "
        "caso cambie de lista (ej.: un rechazo que se reingresó pasa de la 1 a la 2).",
        "4. El resto de las columnas se recalcula desde el sistema: no editarlas.",
        "5. Un caso desaparece de las listas cuando el sistema lo ve resuelto "
        "(entregado, DTE emitido, anulado) y pasa a «4 · Casos cerrados».",
        "6. Marcar «Descartado» NO lo saca de la lista: para que salga, el "
        "movimiento tiene que resolverse en Autoventa/Obuma o anularse.",
    ]
    for k, p in enumerate(pasos, f1 + 1):
        ws.cell(row=k, column=1, value=p)


def _por_vendedor(listas: dict) -> pd.DataFrame:
    partes = []
    for hoja, etq in ((HOJA_SIN_RETOMAR, "Sin retomar"),
                      (HOJA_SIN_CERRAR, "Reingresos sin cerrar"),
                      (HOJA_PEDIDOS, "Pedidos sin gestionar")):
        df = listas[hoja]
        if not df.empty:
            partes.append(df.groupby("Vendedor").size().rename(etq))
    if not partes:
        return pd.DataFrame()
    t = pd.concat(partes, axis=1).fillna(0).astype(int)
    for etq in ("Sin retomar", "Reingresos sin cerrar", "Pedidos sin gestionar"):
        if etq not in t.columns:
            t[etq] = 0
    t = t[["Sin retomar", "Reingresos sin cerrar", "Pedidos sin gestionar"]]
    t["Total"] = t.sum(axis=1)
    t = t.sort_values("Total", ascending=False).reset_index()
    t = t.rename(columns={"index": "Vendedor"})
    tot = {"Vendedor": "TOTAL", **{c: int(t[c].sum()) for c in t.columns[1:]}}
    return pd.concat([t, pd.DataFrame([tot])], ignore_index=True)


def _hoja_casos(wb, titulo: str, df: pd.DataFrame, nota: str,
                manuales: bool = True, validar: bool = True,
                clave: str = ID) -> None:
    """Una lista de casos: encabezado navy, columnas manuales en amarillo,
    lista desplegable en «Estado gestión», filtro y panel congelado."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    ws = wb.create_sheet(titulo[:31])
    ws.cell(row=1, column=1, value=nota).font = Font(italic=True,
                                                     color="64748B")
    fila0 = 3
    if df is None or df.empty:
        ws.cell(row=fila0, column=1, value="Sin casos.")
        return
    cols = list(df.columns)
    if manuales:
        # Las manuales siempre al final, juntas, para que comercial no tenga
        # que buscarlas entre las columnas calculadas.
        auto = [c for c in cols if c not in COLS_MANUALES]
        cols = auto + [c for c in COLS_MANUALES if c in cols]
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=fila0, column=j, value=c)
        es_manual = c in COLS_MANUALES
        x.fill = PatternFill("solid",
                             fgColor=_MANUAL_HDR if es_manual else _HDR_BG)
        x.font = Font(bold=True, color="FFFFFF")
        x.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
    amarillo = PatternFill("solid", fgColor=_MANUAL_BG)
    for i, row in enumerate(df[cols].itertuples(index=False), fila0 + 1):
        for j, (c, v) in enumerate(zip(cols, row), 1):
            x = ws.cell(row=i, column=j, value=_valor(v))
            if isinstance(x.value, date):
                x.number_format = _FMT_FECHA
            if c in COLS_MANUALES and manuales:
                x.fill = amarillo
            if c in ("Acción sugerida", "Lo que dijo el repartidor",
                     "Reintento", "Cómo se cerró",
                     "Comentario / situación del cliente"):
                x.alignment = Alignment(wrap_text=True, vertical="top")
    n = len(df)
    for j, c in enumerate(cols, 1):
        vals = [len(str(v)) for v in df[c].head(200) if v is not None]
        ancho = max([len(str(c))] + vals) + 2
        tope = 48 if c in ("Acción sugerida", "Lo que dijo el repartidor",
                           "Comentario / situación del cliente",
                           "Cómo se cerró", "Reintento") else 34
        ws.column_dimensions[get_column_letter(j)].width = min(max(ancho, 11), tope)
    ws.freeze_panes = ws.cell(row=fila0 + 1, column=3 if clave == ID else 2)
    ws.auto_filter.ref = f"A{fila0}:{get_column_letter(len(cols))}{fila0 + n}"
    if manuales and validar and COL_ESTADO in cols:
        j = cols.index(COL_ESTADO) + 1
        letra = get_column_letter(j)
        dv = DataValidation(type="list",
                            formula1='"' + ",".join(ESTADOS_GESTION) + '"',
                            allow_blank=True)
        dv.error = "Elige un estado de la lista."
        ws.add_data_validation(dv)
        dv.add(f"{letra}{fila0 + 1}:{letra}{fila0 + n + 200}")
    if manuales and "Fecha compromiso" in cols:
        letra = get_column_letter(cols.index("Fecha compromiso") + 1)
        for i in range(fila0 + 1, fila0 + n + 1):
            ws[f"{letra}{i}"].number_format = _FMT_FECHA


def _valor(v):
    """Valor apto para openpyxl."""
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, pd.Timestamp):
        return v.date()
    if hasattr(v, "item"):
        try:
            return v.item()
        except (AttributeError, ValueError):
            pass
    return v
