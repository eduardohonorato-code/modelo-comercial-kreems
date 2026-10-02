"""Excel SIMULADOR de la Propuesta de Comisiones (scorecard de 5 KPIs).

Mismo espíritu que `export_comisiones_sim` (el del modelo de tramos): cada monto
es una fórmula colgada de sus insumos, para que gerencia mueva pisos, pesos y
metas y vea al instante cuánto cambia la comisión, sin tocar el sistema.

Tres hojas:
  · Guía        — qué mide cada indicador, cómo se paga y cómo usar el libro.
  · Parámetros  — pesos, pisos, tope y metas generales LEÍDOS DE LA BASE, más una
                  tabla de sensibilidad (cumple 70/80/90/100%).
  · Simulador   — una fila por vendedor. Amarillo = editable, blanco = fórmula.
                  Trae la comisión que calcula hoy el sistema y la diferencia,
                  que debe dar 0 mientras no se edite nada.

Las fórmulas replican `app/pages/comisiones_v1._calcular` + `_factor`:
logro = real ÷ meta; bajo el piso paga 0; desde el piso paga proporcional al
logro con tope 100%; tasa = suma de los 5 aportes (tope 5%); comisión = tasa ×
venta real.
"""
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from app.export_comisiones_sim import (
    BD, CLP, F_B, F_HD, F_IN, F_IT, F_SEC, F_TIT, F_TXT, FILL_HD, FILL_IN,
    FILL_TOT, MESES, PC0, PC1, _num,
)
from openpyxl.styles import Font, PatternFill

PC2 = "0.00%"
FILL_SIS = PatternFill("solid", fgColor="EDEDED")   # dato del sistema (referencia)
F_SIS = Font(name="Arial", size=10, color="595959")
P = "'Parámetros'!"

# key, nombre, qué mide
KPIS = [
    ("cuota", "Cuota de venta", "Venta real ÷ objetivo de venta del mes"),
    ("nuevos", "Clientes nuevos válidos",
     "Clientes de 1ª compra + reactivados (3+ meses sin comprar) ÷ meta automática"),
    ("cobertura", "Efectividad de cartera",
     "Clientes de la cartera que compraron ÷ (meta % de la temporada × cartera)"),
    ("amplitud", "Amplitud de SKU", "SKUs distintos por cliente ÷ meta de SKUs"),
    ("ruta", "Cobertura de ruta",
     "Visitas ÷ (meta % × agendamientos del reporte de Autoventa)"),
]

# Celdas fijas de la hoja Parámetros (las referencian las fórmulas).
R_MES, R_TEMP, R_TOPE = 4, 5, 6
R_KPI0 = 9                      # filas 9..13 = un KPI cada una
R_EF_V, R_EF_I, R_EF_APL, R_RUTA, R_SKU, R_NV_PCT, R_NV_MIN, R_RE_PCT = range(18, 26)
R_SENS = 29

# Columnas del Simulador: (encabezado, tipo, formato, ancho)
#   in = insumo editable · fx = fórmula · txt = texto · sis = dato del sistema
COLS = [
    ("Vendedor", "txt", None, 30),
    ("Venta real", "in", CLP, 14),
    ("Objetivo de venta", "in", CLP, 14),
    ("Logro cuota", "fx", PC0, 9),
    ("Paga cuota", "fx", PC2, 9),
    ("Nuevos + react.", "in", "0", 9),
    ("Dormidos", "in", "0", 9),
    ("Meta nuevos", "in", "0", 9),
    ("Logro nuevos", "fx", PC0, 9),
    ("Paga nuevos", "fx", PC2, 9),
    ("Clientes que compraron", "in", "0", 10),
    ("Cartera", "in", "0", 9),
    ("Meta efectividad", "fx", "0.0", 10),
    ("Logro efectividad", "fx", PC0, 10),
    ("Paga efectividad", "fx", PC2, 10),
    ("SKUs por cliente", "in", "0.00", 9),
    ("Meta SKUs", "in", "0.0", 9),
    ("Logro SKU", "fx", PC0, 9),
    ("Paga SKU", "fx", PC2, 9),
    ("Visitas", "in", "0", 9),
    ("Agenda-mientos", "in", "0", 10),
    ("Meta visitas", "fx", "0.0", 9),
    ("Logro ruta", "fx", PC0, 9),
    ("Paga ruta", "fx", PC2, 9),
    ("TASA EFECTIVA", "fx", PC2, 10),
    ("COMISIÓN", "fx", CLP, 13),
    ("Comisión en el sistema", "sis", CLP, 13),
    ("Diferencia", "fx", CLP, 12),
]
L = {h: get_column_letter(i) for i, (h, *_r) in enumerate(COLS, start=1)}

# Bloques de encabezado agrupado (fila 3) → (primera col, última col, texto)
GRUPOS = [
    ("Venta real", "Paga cuota", "① CUOTA DE VENTA"),
    ("Nuevos + react.", "Paga nuevos", "② CLIENTES NUEVOS"),
    ("Clientes que compraron", "Paga efectividad", "③ EFECTIVIDAD DE CARTERA"),
    ("SKUs por cliente", "Paga SKU", "④ AMPLITUD DE SKU"),
    ("Visitas", "Paga ruta", "⑤ COBERTURA DE RUTA"),
    ("TASA EFECTIVA", "Diferencia", "RESULTADO"),
]


def _rh(x: str) -> str:
    """Redondeo al par (como round() de Python, que usa el sistema): 2,5→2, 3,5→4."""
    return (f"IF(MOD({x},1)=0.5,IF(MOD(INT({x}),2)=0,INT({x}),INT({x})+1),"
            f"ROUND({x},0))")


def _paga(logro: str, fila_kpi: int) -> str:
    return (f'=IF({logro}="",0,IF({logro}<{P}$D${fila_kpi},0,'
            f'MIN(1,{logro})*{P}$C${fila_kpi}))')


def _hoja_parametros(wb, mes: int, params: dict, pesos: dict):
    ws = wb.create_sheet("Parámetros")
    ws["A1"] = "Parámetros del modelo"
    ws["A1"].font = F_TIT
    ws["A2"] = ("Valores vigentes en el sistema al generar el archivo. Cambia las celdas "
                "amarillas y el Simulador se recalcula. El sistema NO cambia.")
    ws["A2"].font = F_IT

    def fila(r, etiqueta, valor, fmt, editable=True, nota=""):
        ws.cell(r, 1, etiqueta).font = F_B
        c = ws.cell(r, 2, valor)
        c.number_format = fmt
        c.border = BD
        if editable:
            c.font, c.fill = F_IN, FILL_IN
        else:
            c.font = F_TXT
        if nota:
            ws.cell(r, 3, nota).font = F_IT

    fila(R_MES, "Mes del cálculo", mes, "0",
         nota="1–12. Define la temporada de la meta de efectividad.")
    fila(R_TEMP, "Temporada", f'=IF(OR(B{R_MES}>=10,B{R_MES}<=3),"verano","invierno")',
         "@", editable=False, nota="Verano = octubre a marzo.")
    fila(R_TOPE, "Tope de la tasa efectiva", 0.05, PC2,
         nota="Máximo que se paga sobre la venta de cada vendedor.")

    hdr = R_KPI0 - 1
    for j, h in enumerate(["Indicador", "Peso", "% máx sobre venta", "Piso de pago",
                           "Qué mide"], start=1):
        c = ws.cell(hdr, j, h)
        c.font, c.fill, c.border = F_HD, FILL_HD, BD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, (k, nombre, desc) in enumerate(KPIS):
        r = R_KPI0 + i
        ws.cell(r, 1, nombre).font = F_B
        c = ws.cell(r, 2, pesos[k]); c.number_format = PC0
        c.font, c.fill = F_IN, FILL_IN
        c = ws.cell(r, 3, f"=B{r}*$B${R_TOPE}"); c.number_format = PC2; c.font = F_TXT
        c = ws.cell(r, 4, float(params.get(f"umbral_{k}", 0.0))); c.number_format = PC0
        c.font, c.fill = F_IN, FILL_IN
        ws.cell(r, 5, desc).font = F_TXT
        for j in range(1, 6):
            ws.cell(r, j).border = BD
    rt = R_KPI0 + len(KPIS)
    ws.cell(rt, 1, "TOTAL").font = F_B
    for col, fmt in (("B", PC0), ("C", PC2)):
        c = ws[f"{col}{rt}"]
        c.value = f"=SUM({col}{R_KPI0}:{col}{rt - 1})"
        c.number_format, c.font, c.fill = fmt, F_B, FILL_TOT
    ws.cell(rt, 4, "0% = sin piso").font = F_IT
    ws.cell(rt, 5, "Los pesos deberían sumar 100%.").font = F_IT

    ws.cell(R_EF_V - 2, 1, "Metas generales (iguales para todos)").font = F_SEC
    fila(R_EF_V, "Efectividad de cartera · verano", float(params.get("meta_efec_verano", 0.65)),
         PC0, nota="% de la cartera que debe comprar en el mes (oct–mar).")
    fila(R_EF_I, "Efectividad de cartera · invierno",
         float(params.get("meta_efec_invierno", 0.50)), PC0,
         nota="Idem, abr–sep. El helado se mueve ~20 puntos entre temporadas.")
    fila(R_EF_APL, "→ Meta de efectividad aplicada",
         f'=IF(B{R_TEMP}="verano",B{R_EF_V},B{R_EF_I})', PC0, editable=False,
         nota="La que corresponde a la temporada del mes.")
    fila(R_RUTA, "Cobertura de ruta", float(params.get("meta_ruta", 0.90)), PC0,
         nota="% de los agendamientos que debe visitar (industria: sobre 90%).")
    fila(R_SKU, "SKUs distintos por cliente", 5.0, "0.0",
         nota="Meta general; un vendedor con meta manual la trae fija en el Simulador.")
    fila(R_NV_PCT, "Clientes nuevos · % de la cartera", 0.02, PC1,
         nota="Meta automática de nuevos = este % de su cartera…")
    fila(R_NV_MIN, "Clientes nuevos · mínimo", 2, "0", nota="…con este mínimo,")
    fila(R_RE_PCT, "Reactivados · % de sus dormidos", 0.10, PC0,
         nota="…más este % de sus dormidos (mínimo 1 si tiene dormidos).")

    # Sensibilidad
    ws.cell(R_SENS - 1, 1, "Sensibilidad: cuánto paga cada indicador según el cumplimiento"
            ).font = F_SEC
    cab = ["Indicador", "% máx", "Cumple 70%", "Cumple 80%", "Cumple 90%", "Cumple 100%"]
    for j, h in enumerate(cab, start=1):
        c = ws.cell(R_SENS, j, h)
        c.font, c.fill, c.border = F_HD, FILL_HD, BD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, (_k, nombre, _d) in enumerate(KPIS):
        r, src = R_SENS + 1 + i, R_KPI0 + i
        ws.cell(r, 1, nombre).font = F_B
        c = ws.cell(r, 2, f"=C{src}"); c.number_format = PC2; c.font = F_TXT
        for j, x in enumerate([0.7, 0.8, 0.9, 1.0]):
            c = ws.cell(r, 3 + j, f"=IF({x}<$D${src},0,MIN(1,{x})*$C${src})")
            c.number_format, c.font = PC2, F_TXT
        for j in range(1, 7):
            ws.cell(r, j).border = BD
    rs = R_SENS + 1 + len(KPIS)
    ws.cell(rs, 1, "TASA EFECTIVA").font = F_B
    for j in range(2, 7):
        col = get_column_letter(j)
        c = ws.cell(rs, j, f"=MIN($B${R_TOPE},SUM({col}{R_SENS + 1}:{col}{rs - 1}))")
        c.number_format, c.font, c.fill, c.border = PC2, F_B, FILL_TOT, BD
    ws.cell(rs + 1, 1, "Si todos los indicadores cumplen ese %. Bajo el piso, la columna "
            "queda en 0.").font = F_IT

    ws.column_dimensions["A"].width = 36
    for col, w in zip("BCDEF", [14, 16, 13, 13, 13]):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["E"].width = 62
    return ws


def _hoja_guia(wb, anio: int, mes: int):
    ws = wb.create_sheet("Guía")
    ws["A1"] = f"Propuesta de Comisiones · {MESES[mes]} {anio} — simulador para gerencia"
    ws["A1"].font = F_TIT
    bloques = [
        ("Para qué sirve",
         ["Responder «¿qué pasa si…?» sin tocar el sistema: ¿cuánto cuesta bajar el piso al "
          "70%?, ¿y si la meta de efectividad de invierno fuera 45%?, ¿y si la cuota pesa 40%?",
          "La hoja Simulador trae una fila por vendedor con todos los insumos del mes y la "
          "comisión como fórmula. La columna «Diferencia» parte en $0: es la distancia entre lo "
          "que simulas y lo que paga hoy el sistema.",
          "IMPORTANTE: este archivo NO escribe en el sistema. Lo que edites aquí queda aquí."]),
        ("Cómo leerlo",
         ["Celda AMARILLA con número azul = insumo, se puede cambiar.",
          "Celda blanca = fórmula, se recalcula sola. Celda gris = dato del sistema (referencia).",
          "Lo general (pesos, pisos, tope, metas en %) está en la hoja Parámetros y aplica a "
          "todos. Lo de cada vendedor (venta, cartera, visitas…) está en su fila del Simulador."]),
        ("Cómo se calcula (de menos a más)",
         ["1. Comisión = tasa efectiva × venta real del mes. La tasa va de 0% a 5%.",
          "2. La tasa es la suma de 5 indicadores. Cada uno aporta hasta peso × 5% "
          "(cuota 1,50% · nuevos 1,00% · efectividad 1,00% · SKU 0,75% · ruta 0,75%).",
          "3. Logro de cada indicador = lo que hizo ÷ su meta.",
          "4. Piso: si el logro queda bajo el piso, ese indicador paga 0. Desde el piso paga "
          "proporcional al logro (90% → 90% del indicador) y nunca más del 100%.",
          "5. La tasa topa en 5%, pero la comisión en pesos no tiene techo: crece con la venta."]),
        ("Las metas",
         ["Cuota: el objetivo de venta del Panel Gerencia.",
          "Clientes nuevos: automática = máx(2; 2% de la cartera) + máx(1; 10% de sus "
          "dormidos). Si gerencia fijó una meta manual, viene como número fijo.",
          "Efectividad de cartera: % de la cartera según temporada (verano oct–mar / "
          "invierno abr–sep) × cartera.",
          "Amplitud de SKU: meta general de SKUs distintos por cliente (o la manual del vendedor).",
          "Cobertura de ruta: % × agendamientos del reporte de Autoventa. Sin agendamientos "
          "cargados el indicador queda en blanco y no suma."]),
        ("Dónde se cambia de verdad",
         ["Objetivos de venta → Panel Gerencia, editor de objetivos.",
          "Pesos → hoy son fijos en el código; pedir el cambio.",
          "Pisos y metas generales → Comisiones → Propuesta → Configuración → Metas generales y piso.",
          "Metas manuales por vendedor → Configuración → Metas manuales por vendedor.",
          "Visitas y agendamientos → Configuración → Cobertura de ruta (carga mensual)."]),
    ]
    r = 3
    for titulo, lineas in bloques:
        ws.cell(r, 1, titulo).font = F_SEC
        r += 1
        for ln in lineas:
            c = ws.cell(r, 1, ln)
            c.font = F_TXT
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 30
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 118
    return ws


def comisiones_v1_simulador_xlsx(df: pd.DataFrame, anio: int, mes: int,
                                 params: dict, pesos: dict) -> bytes:
    """`df` = salida de comisiones_v1._calcular; `params` = comision_v1_parametro;
    `pesos` = {kpi: peso}."""
    wb = Workbook()
    wb.remove(wb.active)
    _hoja_guia(wb, anio, mes)
    ws = wb.create_sheet("Simulador")
    _hoja_parametros(wb, mes, params, pesos)

    ws["A1"] = f"Simulador · Propuesta de Comisiones · {MESES[mes]} {anio}"
    ws["A1"].font = F_TIT
    ws["A2"] = ("Amarillo = insumo editable · Blanco = fórmula · Gris = lo que calcula hoy el "
                "sistema. Pesos, pisos y metas generales: hoja Parámetros.")
    ws["A2"].font = F_IT

    for desde, hasta, txt in GRUPOS:
        c = ws[f"{L[desde]}3"]
        c.value, c.font = txt, F_B
        c.alignment = Alignment(horizontal="center")
        c.fill = FILL_TOT
        ws.merge_cells(f"{L[desde]}3:{L[hasta]}3")

    hdr = 4
    for j, (h, *_r) in enumerate(COLS, start=1):
        c = ws.cell(hdr, j, h)
        c.font, c.fill, c.border = F_HD, FILL_HD, BD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[hdr].height = 40

    k = {key: R_KPI0 + i for i, (key, *_r) in enumerate(KPIS)}
    r = hdr + 1
    for _, x in df.sort_values("nombre_canonico").iterrows():
        c_ = {h: f"{L[h]}{r}" for h in L}
        ov_nv = x.get("ov_meta_nuevos_react")
        ov_sku = x.get("ov_meta_lineas")
        cart, dorm = c_["Cartera"], c_["Dormidos"]
        meta_nv_auto = (f"=MAX({P}$B${R_NV_MIN},{_rh(f'{P}$B${R_NV_PCT}*{cart}')})"
                        f"+IF({dorm}>0,MAX(1,{_rh(f'{P}$B${R_RE_PCT}*{dorm}')}),0)")
        vals = {
            "Vendedor": x["nombre_canonico"],
            "Venta real": _num(x.get("fact_nc")),
            "Objetivo de venta": _num(x.get("cuota_meta")),
            "Nuevos + react.": _num(x.get("nuevos_real")),
            "Dormidos": _num(x.get("dormidos")),
            "Meta nuevos": (_num(ov_nv) if ov_nv is not None and pd.notna(ov_nv)
                            else meta_nv_auto),
            "Clientes que compraron": _num(x.get("cobertura_real")),
            "Cartera": _num(x.get("cartera")),
            "SKUs por cliente": _num(x.get("amplitud_real")),
            "Meta SKUs": (_num(ov_sku) if ov_sku is not None and pd.notna(ov_sku)
                          else f"={P}$B${R_SKU}"),
            "Visitas": _num(x.get("visitas")),
            "Agenda-mientos": _num(x.get("agendamientos")),
            "Comisión en el sistema": round(_num(x.get("comision_total"))),
        }

        def lg(real, meta):
            return f'=IF(N({c_[meta]})=0,"",{c_[real]}/{c_[meta]})'

        fx = {
            "Logro cuota": lg("Venta real", "Objetivo de venta"),
            "Paga cuota": _paga(c_["Logro cuota"], k["cuota"]),
            "Logro nuevos": lg("Nuevos + react.", "Meta nuevos"),
            "Paga nuevos": _paga(c_["Logro nuevos"], k["nuevos"]),
            "Meta efectividad": f"={cart}*{P}$B${R_EF_APL}",
            "Logro efectividad": lg("Clientes que compraron", "Meta efectividad"),
            "Paga efectividad": _paga(c_["Logro efectividad"], k["cobertura"]),
            "Logro SKU": lg("SKUs por cliente", "Meta SKUs"),
            "Paga SKU": _paga(c_["Logro SKU"], k["amplitud"]),
            "Meta visitas": f"={c_['Agenda-mientos']}*{P}$B${R_RUTA}",
            "Logro ruta": lg("Visitas", "Meta visitas"),
            "Paga ruta": _paga(c_["Logro ruta"], k["ruta"]),
            "TASA EFECTIVA": (f"=MIN({P}$B${R_TOPE},{c_['Paga cuota']}+{c_['Paga nuevos']}"
                              f"+{c_['Paga efectividad']}+{c_['Paga SKU']}+{c_['Paga ruta']})"),
            "COMISIÓN": f"=ROUND({c_['TASA EFECTIVA']}*{c_['Venta real']},0)",
            "Diferencia": f"={c_['COMISIÓN']}-{c_['Comisión en el sistema']}",
        }
        for j, (h, tipo, fmt, _w) in enumerate(COLS, start=1):
            cel = ws.cell(r, j, fx[h] if tipo == "fx" else vals.get(h))
            cel.border = BD
            if fmt:
                cel.number_format = fmt
            if tipo == "in":
                cel.font, cel.fill = F_IN, FILL_IN
            elif tipo == "sis":
                cel.font, cel.fill = F_SIS, FILL_SIS
            elif h in ("TASA EFECTIVA", "COMISIÓN"):
                cel.font, cel.fill = F_B, FILL_TOT
            else:
                cel.font = F_TXT
        r += 1

    ult = r - 1
    ws.cell(r, 1, "TOTAL").font = F_B
    for h in ("Venta real", "Objetivo de venta", "COMISIÓN", "Comisión en el sistema",
              "Diferencia"):
        cel = ws[f"{L[h]}{r}"]
        cel.value = f"=SUM({L[h]}{hdr + 1}:{L[h]}{ult})"
        cel.number_format, cel.font, cel.fill, cel.border = CLP, F_B, FILL_TOT, BD
    cel = ws[f"{L['TASA EFECTIVA']}{r}"]
    cel.value = f'=IF({L["Venta real"]}{r}=0,0,{L["COMISIÓN"]}{r}/{L["Venta real"]}{r})'
    cel.number_format, cel.font, cel.fill, cel.border = PC2, F_B, FILL_TOT, BD

    for j, (_h, _t, _f, w) in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(hdr + 1, 2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
